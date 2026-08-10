from pathlib import Path


ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def extract_text(file_path: Path) -> str:
    """Extract plain text from .txt, .pdf, or .docx."""
    suffix = file_path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Allowed: {sorted(ALLOWED_EXTENSIONS)}"
        )

    if suffix == ".txt":
        return _extract_txt(file_path)
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    return _extract_docx(file_path)


def _extract_txt(file_path: Path) -> str:
    raw = file_path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")

    text = text.strip()
    if not text:
        raise ValueError("Document contains no extractable text")
    return text


def _extract_pdf(file_path: Path) -> str:
    import pymupdf

    parts: list[str] = []
    with pymupdf.open(file_path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("PDF contains no extractable text")
    return text


def _extract_docx(file_path: Path) -> str:
    from docx import Document

    doc = Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    parts.append(cell_text)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("DOCX contains no extractable text")
    return text
